"""Builds KrishiRakshak_Presentation.pptx following the Joy University project template."""
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SHOTS = os.path.join(HERE, "screenshots")
ASSETS = os.path.join(HERE, "assets")
PLOTS = os.path.join(ROOT, "ml", "plots")
M = json.load(open(os.path.join(ROOT, "ml", "model", "metrics.json")))
S = json.load(open(os.path.join(HERE, "stats.json"), encoding="utf-8"))
import sys
sys.path.insert(0, ROOT)
import kb as _kb  # noqa: E402
NA = len(_kb.ACTIONS)

GREEN = RGBColor(0x14, 0x5A, 0x32)
LGREEN = RGBColor(0xE3, 0xF2, 0xE8)
ACCENT = RGBColor(0x7C, 0xB3, 0x42)
ORANGE = RGBColor(0xE6, 0x51, 0x00)
LORANGE = RGBColor(0xFF, 0xF3, 0xE0)
BLUE = RGBColor(0x44, 0x72, 0xC4)
RED = RGBColor(0xFF, 0x00, 0x00)
INK = RGBColor(0x1D, 0x27, 0x21)
MUTED = RGBColor(0x5D, 0x6B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def text(slide, x, y, w, h, s, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, font
    return tb


def bullets(slide, x, y, w, h, items, size=15, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        sub = isinstance(it, tuple)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = ("–  " if sub else "•  ") + (it[0] if sub else it)
        r.font.size = Pt(size - 2 if sub else size)
        r.font.color.rgb = MUTED if sub else INK
        r.font.name = "Calibri"
        if sub:
            p.level = 1
    return tb


def rich_bullets(slide, x, y, w, h, items, size=14, gap=5):
    """items: list of (bold_head, rest)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (hd, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "•  " + hd; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = GREEN; r.font.name = "Calibri"
        r = p.add_run(); r.text = " " + rest; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Calibri"
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def box(slide, x, y, w, h, s, fill=LGREEN, line=GREEN, size=11, bold=False, color=INK):
    b = rect(slide, x, y, w, h, fill, line)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = s.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
        r.font.bold = bold or (i == 0 and len(lines) > 1)
    return b


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    pic = slide.shapes.add_picture(path, x, y, w, h)
    if border:
        pic.line.color.rgb = RGBColor(0xDD, 0xE5, 0xDF); pic.line.width = Pt(0.75)
    return pic


def fit_picture(slide, path, x, y, maxw, maxh, border=True, center=True):
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(maxw / iw, maxh / ih)
    w, h = int(iw * scale), int(ih * scale)
    if center:
        x, y = x + (maxw - w) // 2, y + (maxh - h) // 2
    return picture(slide, path, x, y, w, h, border)


def crop_picture(slide, path, x, y, w, h, keep_top=1.0):
    """Place a (tall) screenshot cropped to the box aspect ratio, keeping the top part."""
    from PIL import Image
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, x, y, w, h)
    target = h / w
    img_ratio = ih / iw
    if img_ratio > target:
        pic.crop_bottom = 1 - target / img_ratio
    else:
        cut = (1 - img_ratio / target) / 2
        pic.crop_left = pic.crop_right = cut
    pic.line.color.rgb = RGBColor(0xDD, 0xE5, 0xDF); pic.line.width = Pt(0.75)
    return pic


def content_slide(title, section, num):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, Inches(0.12), GREEN, shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.5), Inches(0.35), Inches(7.6), Inches(0.8), title, size=28, color=INK)
    text(s, Inches(7.6), Inches(0.42), Inches(2.0), Inches(0.35), section.upper(), size=10, bold=True, color=ACCENT, align=PP_ALIGN.RIGHT)
    rect(s, Inches(0.5), Inches(1.12), Inches(1.2), Inches(0.05), ACCENT, shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.5), Inches(7.05), Inches(6), Inches(0.3), "KrishiRakshak · Crop disease & pest early-warning system", size=9, color=MUTED)
    text(s, Inches(8.8), Inches(7.05), Inches(0.8), Inches(0.3), str(num), size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


pct = lambda v: f"{v * 100:.2f}%"
n = 1

# ------------------------------------------------------------------ 1. Title (template layout)
s = prs.slides.add_slide(BLANK)
fit_picture(s, os.path.join(ASSETS, "tpl_img0.jpeg"), Inches(1.6), Inches(0.15), Inches(6.8), Inches(1.9), border=False)
text(s, Inches(1), Inches(2.05), Inches(8), Inches(0.4), "Course: Operating system", size=20, color=RED, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(2.45), Inches(8), Inches(0.4), "Subject code : 24BTDS145", size=20, color=RED, align=PP_ALIGN.CENTER)
text(s, Inches(0.5), Inches(3.0), Inches(9), Inches(1.2),
     ["KrishiRakshak: AI-Based Early Detection and", "Management of Crop Diseases & Pest Infestations"],
     size=26, color=INK, align=PP_ALIGN.CENTER)
rows = [("Member PRN No", "Name")] + [("24BTDSXXX", "[Member name]")] * 4
tbl = s.shapes.add_table(len(rows), 2, Inches(2.4), Inches(4.45), Inches(5.2), Inches(2.4)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(1.6), Inches(3.6)
for r, (a, b) in enumerate(rows):
    for c, v in enumerate((a, b)):
        cell = tbl.cell(r, c)
        cell.text = v
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(12)
        para.runs[0].font.bold = r == 0
        para.runs[0].font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if r == 0 else (RGBColor(0xCF, 0xD5, 0xEA) if r % 2 else RGBColor(0xE9, 0xEB, 0xF5))

# ------------------------------------------------------------------ 2. Introduction – context
n += 1
s = content_slide("Introduction", "Introduction", n)
text(s, Inches(0.5), Inches(1.35), Inches(9), Inches(0.4), "Context and background", size=18, bold=True, color=GREEN)
bullets(s, Inches(0.5), Inches(1.85), Inches(5.3), Inches(4.8), [
    "Plant pests and diseases destroy up to 40% of global crop production every year (FAO).",
    "Farmers usually notice a disease only after visible damage has spread across the field.",
    "Extension officers cover very large areas; lab diagnosis and expert advice are slow to reach villages.",
    "Weather, crop stage, variety, soil and local pest history drive outbreaks — but are rarely combined into farm-level alerts.",
    "Wrong or late diagnosis leads to excess pesticide use, higher cost, residues and yield loss.",
], size=14)
for i, (v, l) in enumerate([("40%", "of crop production lost to pests & diseases (FAO)"),
                            ("38", "crop disease & health classes recognised by our AI"),
                            ("6", "languages: English, हिन्दी, తెలుగు, தமிழ், ಕನ್ನಡ, मराठी")]):
    y = Inches(1.9 + i * 1.6)
    rect(s, Inches(6.2), y, Inches(3.3), Inches(1.35), LGREEN)
    text(s, Inches(6.35), y + Inches(0.08), Inches(3.0), Inches(0.6), v, size=30, bold=True, color=GREEN)
    text(s, Inches(6.35), y + Inches(0.72), Inches(3.0), Inches(0.6), l, size=12, color=MUTED)

# ------------------------------------------------------------------ 3. Problem statement
n += 1
s = content_slide("Problem Statement", "Introduction", n)
text(s, Inches(0.5), Inches(1.35), Inches(9), Inches(0.5), "Early detection and management of crop diseases and pest infestations", size=18, bold=True, color=GREEN)
text(s, Inches(0.5), Inches(1.85), Inches(9), Inches(0.5), "Why is this project necessary?", size=14, color=MUTED)
probs = [("Late detection", "Symptoms are recognised only after the damage spreads."),
         ("Limited expert reach", "Few extension staff and labs for many farmers."),
         ("Scattered risk signals", "Weather, crop stage, soil and pest history are not combined."),
         ("Wrong treatment", "Misdiagnosis → excess or wrong pesticide, cost, residues."),
         ("No surveillance picture", "Officials lack real-time outbreak maps and data."),
         ("Language barrier", "Advice is rarely available in the farmer's language.")]
for i, (h, d) in enumerate(probs):
    col, row = i % 2, i // 2
    x, y = Inches(0.5 + col * 4.6), Inches(2.45 + row * 1.45)
    rect(s, x, y, Inches(4.4), Inches(1.25), LORANGE, RGBColor(0xFF, 0xCC, 0x80))
    text(s, x + Inches(0.15), y + Inches(0.1), Inches(4.1), Inches(0.4), h, size=15, bold=True, color=ORANGE)
    text(s, x + Inches(0.15), y + Inches(0.52), Inches(4.1), Inches(0.7), d, size=13)

# ------------------------------------------------------------------ 4. Objectives
n += 1
s = content_slide("Objectives and Key Goals", "Introduction", n)
objs = [("Detect", "Identify diseases from a leaf photo with a CNN model (>95% target accuracy)."),
        ("Forecast", "7-day weather-based disease & pest risk for each farm."),
        ("Monitor", "Pest-trap counts and IoT sensors with economic-threshold (ETL) alerts."),
        ("Map", "Automatic geospatial hotspot detection and alerts to nearby farms."),
        ("Validate", "Expert confirmation / correction of every AI diagnosis; lab referral."),
        ("Advise", "Integrated Pest Management (IPM) advice with safe pesticide use, in 6 languages with audio."),
        ("Follow up", "3/7-day follow-up checks; escalate cases that are not improving."),
        ("Learn", "Retrain the model from field-confirmed images (continual learning)."),
        ("Inform", "Dashboards with surveillance KPIs for agriculture officials.")]
for i, (h, d) in enumerate(objs):
    col, row = i % 3, i // 3
    x, y = Inches(0.5 + col * 3.05), Inches(1.45 + row * 1.8)
    rect(s, x, y, Inches(2.85), Inches(1.62), LGREEN)
    text(s, x + Inches(0.12), y + Inches(0.08), Inches(2.6), Inches(0.4), f"{i + 1}. {h}", size=15, bold=True, color=GREEN)
    text(s, x + Inches(0.12), y + Inches(0.5), Inches(2.65), Inches(1.1), d, size=12)

# ------------------------------------------------------------------ 5. Methodology – approach / architecture
n += 1
s = content_slide("Methodology: Approach", "Methodology", n)
text(s, Inches(0.5), Inches(1.3), Inches(9), Inches(0.4),
     "Three-tier web system: browser front-end → Flask server with six intelligence modules → data & external services",
     size=13, color=MUTED)
users = ["Farmer\nscan · farms · traps · alerts", "Extension officer\nvalidation · referral · map", "Agriculture official\ndashboard · KPIs · model"]
for i, u in enumerate(users):
    box(s, Inches(0.5 + i * 3.05), Inches(1.8), Inches(2.85), Inches(0.75), u, fill=WHITE, size=11)
    arrow(s, Inches(1.92 + i * 3.05), Inches(2.55), Inches(1.92 + i * 3.05), Inches(2.8))
box(s, Inches(0.5), Inches(2.8), Inches(9.0), Inches(0.5), "Chrome front-end: HTML5 · CSS3 · JavaScript · Jinja2 · Leaflet.js · Chart.js · Web Speech API",
    fill=RGBColor(0xE3, 0xF2, 0xFD), line=RGBColor(0x15, 0x65, 0xC0), size=11)
arrow(s, Inches(5), Inches(3.3), Inches(5), Inches(3.55))
mods = ["AI diagnosis\nMobileNetV2", "Risk forecast\nweather models", "Trap / IoT\nETL engine", "Hotspots\nDBSCAN", "IPM advisory\n6 languages", "Validation &\nlearning"]
for i, mname in enumerate(mods):
    box(s, Inches(0.5 + i * 1.52), Inches(3.55), Inches(1.4), Inches(1.0), mname, size=10)
box(s, Inches(0.5), Inches(4.7), Inches(9.0), Inches(0.42), "Flask (Python) · role-based auth · alert engine · follow-up scheduler · REST APIs",
    fill=RGBColor(0xEC, 0xEF, 0xF1), line=RGBColor(0x78, 0x90, 0x9C), size=11)
data = ["SQLite\ndatabase", "Model store\nKeras heads v1..vN", "Knowledge base\n39 diseases · 7 pests", "Open-Meteo · OSM\nIoT devices"]
for i, dname in enumerate(data):
    arrow(s, Inches(1.62 + i * 2.28), Inches(5.12), Inches(1.62 + i * 2.28), Inches(5.4))
    box(s, Inches(0.5 + i * 2.28), Inches(5.4), Inches(2.15), Inches(0.9), dname, fill=LORANGE, line=ORANGE, size=11)
text(s, Inches(0.5), Inches(6.45), Inches(9), Inches(0.5),
     "Approach: combine computer vision (what the farmer sees) with agro-meteorological risk models (what is likely next) and human expert validation (trust).",
     size=12, color=GREEN, bold=True)

# ------------------------------------------------------------------ 6. Tools
n += 1
s = content_slide("Methodology: Tools & Technologies", "Methodology", n)
rows = [("Layer", "Technology", "Purpose"),
        ("Front-end", "HTML5, CSS3, JavaScript, Jinja2", "Responsive pages that run in Chrome (desktop & mobile)"),
        ("Maps & charts", "Leaflet.js + OpenStreetMap, Leaflet.heat, Chart.js", "Hotspot map, heatmap, dashboards, weather chart"),
        ("Back-end", "Python 3.12, Flask 3", "Routes, auth, REST APIs, alert & follow-up engine"),
        ("Deep learning", "TensorFlow 2.18 / Keras, MobileNetV2", "Leaf image classification (transfer learning)"),
        ("ML utilities", "scikit-learn, NumPy, Pillow, Matplotlib", "DBSCAN hotspots, metrics, image quality, plots"),
        ("Database", "SQLite", "Users, farms, reports, traps, sensors, alerts"),
        ("Weather", "Open-Meteo forecast API (free)", "Hourly temperature, humidity, rain for 10 days"),
        ("Voice", "Web Speech API (browser)", "Read advisories aloud in the farmer's language"),
        ("Hardware (sim.)", "IoT weather station / smart trap via /api/sensor", "Sensor & trap-count input")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.4), Inches(9), Inches(5.3)).table
for c, wdt in enumerate((1.6, 3.4, 4.0)):
    tbl.columns[c].width = Inches(wdt)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12 if r == 0 else 11)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN if r == 0 else (LGREEN if r % 2 else WHITE)

# ------------------------------------------------------------------ 7. Data collection
n += 1
s = content_slide("Methodology: Data Collection", "Methodology", n)
rich_bullets(s, Inches(0.5), Inches(1.4), Inches(4.9), Inches(5.4), [
    ("Image dataset:", f"PlantVillage (Mendeley Data, without augmentation) — {M['n_images']:,} leaf images, "
                       f"{M['n_classes']} classes (38 crop/disease classes of 14 crops + background)."),
    ("Split:", f"stratified 70 / 15 / 15 → {M['split']['train']:,} train, {M['split']['val']:,} validation, {M['split']['test']:,} test images."),
    ("Weather:", "Open-Meteo hourly forecast (past 3 + next 7 days) for each farm's GPS location."),
    ("Knowledge base:", "39 disease and 7 pest profiles — pathogen, symptoms, IPM actions, doses, pre-harvest intervals, ETLs — compiled from IPM package-of-practice literature."),
    ("Field data:", f"farm profiles, trap counts, sensor readings and reports entered through the app. Demo database: {S['farms']} farms, {S['reports']} reports (synthetic, for evaluation)."),
], size=13)
fit_picture(s, os.path.join(PLOTS, "class_distribution.png"), Inches(5.55), Inches(1.35), Inches(4.0), Inches(5.55))

# ------------------------------------------------------------------ 8. Process – workflow
n += 1
s = content_slide("Process: Workflow", "Process", n)
steps = ["1. Capture\nleaf photo", "2. Quality\ncheck", "3. AI\ndiagnosis", "4. Risk\ncontext", "5. IPM\nadvisory"]
for i, st in enumerate(steps):
    x = Inches(0.5 + i * 1.84)
    box(s, x, Inches(1.5), Inches(1.6), Inches(1.0), st, size=12)
    if i < 4:
        arrow(s, x + Inches(1.6), Inches(2.0), x + Inches(1.84), Inches(2.0))
steps2 = ["10. Dashboard\nfor officials", "9. Continual\nlearning", "8. Follow-up\n3 / 7 days", "7. Hotspot\nalerts", "6. Expert\nvalidation"]
for i, st in enumerate(steps2):
    x = Inches(0.5 + i * 1.84)
    box(s, x, Inches(3.1), Inches(1.6), Inches(1.0), st, fill=LORANGE, line=ORANGE, size=12)
    if i < 4:
        arrow(s, x + Inches(1.84), Inches(3.6), x + Inches(1.6), Inches(3.6))
arrow(s, Inches(8.16), Inches(2.5), Inches(8.16), Inches(3.1))
rich_bullets(s, Inches(0.5), Inches(4.4), Inches(9), Inches(2.6), [
    ("Farmer:", "registers farm (GPS, crop, variety, sowing date, soil, irrigation) → gets daily risk alerts → scans a sick leaf → receives advice + audio."),
    ("System:", "stores the image embedding, schedules a follow-up, refers to the nearest lab when confidence < 60% or disease is notifiable/viral."),
    ("Officer:", "confirms or corrects the diagnosis and sets severity → farmer notified → hotspots recomputed → confirmed image added to training data."),
], size=13)

# ------------------------------------------------------------------ 9. Implementation – AI model
n += 1
s = content_slide("Process: AI Diagnosis Model", "Process", n)
chain = ["Leaf image\n224×224×3", "MobileNetV2\n(ImageNet, frozen)", "1280-d\nfeature", "Dropout 0.3 →\nDense 256 ReLU", f"Softmax\n{M['n_classes']} classes"]
for i, c in enumerate(chain):
    x = Inches(0.5 + i * 1.84)
    box(s, x, Inches(1.45), Inches(1.6), Inches(0.95), c, fill=LGREEN if i != 1 else RGBColor(0xE3, 0xF2, 0xFD), size=11)
    if i < 4:
        arrow(s, x + Inches(1.6), Inches(1.92), x + Inches(1.84), Inches(1.92))
rich_bullets(s, Inches(0.5), Inches(2.7), Inches(9), Inches(4.2), [
    ("Transfer learning:", "ImageNet-pretrained MobileNetV2 is used as a frozen feature extractor — light enough for a CPU-only laptop and mobile deployment."),
    ("Two-stage training:", f"features of all {M['n_images']:,} images are computed once (~40 min on CPU); the classification head then trains in {M['head_train_seconds']:.0f} s."),
    ("Imbalance:", "balanced class weights (classes range from 152 to 5,507 images); early stopping on validation accuracy; learning-rate decay on plateau."),
    ("Crop-aware inference:", "if the farmer selects the crop, probabilities are re-normalised over that crop's classes; a strong mismatch is flagged."),
    ("Image quality gate:", "Laplacian-variance blur check and exposure check ask the farmer to retake poor photos."),
    ("Continual learning:", "each expert-confirmed image's embedding is stored; retraining fine-tunes the head (field samples ×5) and activates the new version only if held-out test accuracy does not drop > 0.5 pt."),
], size=13)

# ------------------------------------------------------------------ 10. Implementation – risk / hotspots / advisory
n += 1
s = content_slide("Process: Risk, Hotspots & Advisory Logic", "Process", n)
rect(s, Inches(0.5), Inches(1.4), Inches(9), Inches(0.85), LGREEN)
text(s, Inches(0.65), Inches(1.47), Inches(8.7), Inches(0.75),
     ["Risk score = 85 × weather favourability × crop-stage factor × variety resistance × drainage/irrigation factor",
      "                   + 6 × confirmed cases within 10 km in 21 days (max +25)  →  Low / Moderate / High / Severe"],
     size=12, bold=True, color=GREEN)
rich_bullets(s, Inches(0.5), Inches(2.45), Inches(9), Inches(4.5), [
    ("Weather models:", "12 disease models, e.g. late blight — Hutton criteria (Tmin ≥ 10 °C and ≥ 6 h RH ≥ 90%); apple scab — Mills wet-hour table; spider mites — hot & dry spells."),
    ("Crop factors:", "seedling 0.85 · vegetative 1.0 · flowering 1.15 · maturity 0.75; resistant variety 0.5, moderately resistant 0.75; poor drainage ×1.15, sprinkler ×1.1."),
    ("Pest traps & sensors:", "counts normalised per trap per day/week and compared with ETL (e.g. pink bollworm 8 moths/trap/night, BPH 10/hill); IoT leaf-wetness raises humid-disease risk."),
    ("Hotspots:", "DBSCAN with haversine distance (eps 5 km, ≥ 3 cases, last 30 days) per disease; farms of the same crop within radius + 5 km are alerted."),
    ("IPM advisory:", f"cultural → biological → chemical (last resort, dose + pre-harvest interval) → safety → referral → follow-up; composed from {NA} translated action phrases."),
], size=13)

# ------------------------------------------------------------------ 11. Challenges
n += 1
s = content_slide("Process: Challenges & Adaptations", "Process", n)
ch = [("No GPU available", "Frozen-backbone feature caching: one-time CPU feature extraction, then seconds-long head training and retraining."),
      ("Class imbalance (152–5,507 images)", "Balanced class weights and per-class F1 reporting instead of accuracy alone."),
      ("Lab images vs field photos", "Image-quality gate, crop-aware re-normalisation, low-confidence referral and expert validation loop."),
      ("Translating advice to 6 languages", f"Phrase library of {NA} IPM actions with templated doses — new diseases reuse translated phrases."),
      ("Weather API may be offline", "Hourly caching + seasonal offline fallback clearly labelled in the UI."),
      ("Bad labels could harm the model", "New model versions are activated only if held-out accuracy does not drop.")]
for i, (h, d) in enumerate(ch):
    y = Inches(1.45 + i * 0.9)
    rect(s, Inches(0.5), y, Inches(3.3), Inches(0.78), LORANGE, RGBColor(0xFF, 0xCC, 0x80))
    text(s, Inches(0.6), y, Inches(3.1), Inches(0.78), h, size=13, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, Inches(3.85), y + Inches(0.39), Inches(4.15), y + Inches(0.39))
    rect(s, Inches(4.2), y, Inches(5.3), Inches(0.78), LGREEN)
    text(s, Inches(4.3), y, Inches(5.1), Inches(0.78), d, size=12, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 12. Results – model metrics
n += 1
s = content_slide("Results: Model Performance", "Results", n)
kpis = [(pct(M["test_accuracy"]), "Test accuracy"), (pct(M["test_top3_accuracy"]), "Top-3 accuracy"),
        (f"{M['macro_f1']:.3f}", "Macro F1-score"), (f"{M['split']['test']:,}", "Unseen test images")]
for i, (v, l) in enumerate(kpis):
    x = Inches(0.5 + i * 2.28)
    rect(s, x, Inches(1.4), Inches(2.1), Inches(1.1), LGREEN)
    text(s, x, Inches(1.45), Inches(2.1), Inches(0.6), v, size=24, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    text(s, x, Inches(2.05), Inches(2.1), Inches(0.4), l, size=12, color=MUTED, align=PP_ALIGN.CENTER)
fit_picture(s, os.path.join(PLOTS, "training_curves.png"), Inches(0.5), Inches(2.7), Inches(5.6), Inches(2.4))
worst = sorted(M["per_class"].items(), key=lambda kv: kv[1]["f1-score"])[:5]
rows = [("Weakest classes", "F1")] + [(_kb.full_name(l, "en"), f"{v['f1-score']:.3f}") for l, v in worst]
tbl = s.shapes.add_table(len(rows), 2, Inches(6.3), Inches(2.75), Inches(3.2), Inches(2.3)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(2.5), Inches(0.7)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c); cell.text = v
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(10); run.font.bold = r == 0; run.font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid(); cell.fill.fore_color.rgb = GREEN if r == 0 else (LGREEN if r % 2 else WHITE)
bullets(s, Inches(0.5), Inches(5.3), Inches(9), Inches(1.6), [
    f"Validation curves track training curves closely — no over-fitting; early stopping after {M['epochs_run']} epochs.",
    "Most errors are between visually similar leaf-spot diseases of the same crop (e.g. maize gray leaf spot vs northern leaf blight).",
    "Low-confidence predictions are routed to experts, so residual errors are caught before advice is finalised.",
], size=12, gap=3)

# ------------------------------------------------------------------ 13. Results – confusion matrix
n += 1
s = content_slide("Results: Confusion Matrix", "Results", n)
fit_picture(s, os.path.join(PLOTS, "confusion_matrix.png"), Inches(0.4), Inches(1.3), Inches(6.2), Inches(5.65))
rich_bullets(s, Inches(6.75), Inches(1.5), Inches(2.85), Inches(5.4), [
    ("Strong diagonal:", "almost every class is recognised correctly on unseen test images."),
    ("Background class:", "non-leaf photos are detected and the farmer is asked to retake."),
    ("Confusions:", "mainly within the same crop — the crop-aware mode and the top-3 list help the expert decide."),
    ("Benchmark:", "Mohanty et al. (2016) reached 99.35% with fully fine-tuned GoogLeNet on PlantVillage; our frozen MobileNetV2 is far lighter and trains on a laptop CPU."),
], size=12)

# ------------------------------------------------------------------ 14. Results – farmer app
n += 1
s = content_slide("Results: Farmer Application", "Results", n)
crop_picture(s, os.path.join(SHOTS, "02_farmer_home.png"), Inches(0.4), Inches(1.35), Inches(4.5), Inches(3.2))
crop_picture(s, os.path.join(SHOTS, "05_report.png"), Inches(5.1), Inches(1.35), Inches(4.5), Inches(3.2))
text(s, Inches(0.4), Inches(4.58), Inches(4.5), Inches(0.3), "Farmer home: 7-day risk per farm, live weather, alerts", size=10, color=MUTED, align=PP_ALIGN.CENTER)
text(s, Inches(5.1), Inches(4.58), Inches(4.5), Inches(0.3), "Diagnosis: top-3 predictions, weather link, IPM advisory", size=10, color=MUTED, align=PP_ALIGN.CENTER)
crop_picture(s, os.path.join(SHOTS, "06_report_kannada.png"), Inches(0.4), Inches(4.95), Inches(3.0), Inches(1.95))
crop_picture(s, os.path.join(SHOTS, "07_advisory_tamil.png"), Inches(3.55), Inches(4.95), Inches(3.0), Inches(1.95))
crop_picture(s, os.path.join(SHOTS, "03_farm_risk.png"), Inches(6.7), Inches(4.95), Inches(2.9), Inches(1.95))

# ------------------------------------------------------------------ 15. Results – officials
n += 1
s = content_slide("Results: Officer & Official Views", "Results", n)
crop_picture(s, os.path.join(SHOTS, "12_dashboard.png"), Inches(0.4), Inches(1.35), Inches(4.5), Inches(3.3))
crop_picture(s, os.path.join(SHOTS, "11_map.png"), Inches(5.1), Inches(1.35), Inches(4.5), Inches(3.3))
text(s, Inches(0.4), Inches(4.68), Inches(4.5), Inches(0.3), "Surveillance dashboard (KPIs & charts)", size=10, color=MUTED, align=PP_ALIGN.CENTER)
text(s, Inches(5.1), Inches(4.68), Inches(4.5), Inches(0.3), "Hotspot map: DBSCAN clusters + heatmap", size=10, color=MUTED, align=PP_ALIGN.CENTER)
k = S["kpi"]
kp = [(f"{k['agreement']}%", "AI–expert agreement"), (f"{k['avg_response_h']} h", "avg. validation time"),
      (str(S["hotspots"]), "hotspots detected"), (f"{k['no_chem_pct']}%", "cases managed without chemicals")]
for i, (v, l) in enumerate(kp):
    x = Inches(0.4 + i * 2.32)
    rect(s, x, Inches(5.15), Inches(2.15), Inches(1.0), LGREEN)
    text(s, x, Inches(5.2), Inches(2.15), Inches(0.5), v, size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    text(s, x, Inches(5.7), Inches(2.15), Inches(0.4), l, size=11, color=MUTED, align=PP_ALIGN.CENTER)
text(s, Inches(0.4), Inches(6.3), Inches(9.2), Inches(0.6),
     f"Measured on the demo database ({S['reports']} reports over 90 days; AI predictions are real model outputs on held-out validation images, field context is simulated).",
     size=10, color=MUTED)

# ------------------------------------------------------------------ 16. Comparison
n += 1
s = content_slide("Results: Comparison with Existing Approaches", "Results", n)
rows = [("Capability", "Manual extension", "Photo-ID apps", "KrishiRakshak"),
        ("Image-based diagnosis", "✗ (visit needed)", "✓", "✓ 38 classes + top-3"),
        ("Weather risk forecast per farm", "✗", "✗ / generic", "✓ 12 models, 7 days"),
        ("Crop stage / variety / soil", "partly", "✗", "✓"),
        ("Pest traps & IoT sensors", "manual", "✗", "✓ ETL alerts + API"),
        ("Hotspot mapping", "✗", "✗", "✓ DBSCAN + heatmap"),
        ("Expert validation loop", "✓ (slow)", "rare", "✓ prioritised queue"),
        ("Learns from field confirmations", "—", "✗", "✓ versioned retraining"),
        ("Multilingual + audio advice", "partly", "partly", "✓ 6 languages"),
        ("Officials' dashboard", "✗", "✗", "✓ KPIs & charts")]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.4), Inches(9), Inches(5.2)).table
for c, wdt in enumerate((3.1, 1.8, 1.8, 2.3)):
    tbl.columns[c].width = Inches(wdt)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c); cell.text = v
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12); run.font.bold = r == 0 or c == 3
        run.font.color.rgb = WHITE if r == 0 else (GREEN if c == 3 else INK)
        cell.fill.solid(); cell.fill.fore_color.rgb = GREEN if r == 0 else (LGREEN if r % 2 else WHITE)

# ------------------------------------------------------------------ 17. Discussion
n += 1
s = content_slide("Discussion", "Discussion", n)
text(s, Inches(0.5), Inches(1.35), Inches(4.4), Inches(0.4), "Interpretation of results", size=17, bold=True, color=GREEN)
bullets(s, Inches(0.5), Inches(1.85), Inches(4.4), Inches(5), [
    f"A lightweight CNN reaches {pct(M['test_accuracy'])} accuracy — enough for a first-line screening tool.",
    "Combining the photo diagnosis with weather risk tells the farmer not only what the problem is but how urgent it is.",
    "The expert loop turns AI into a decision-support tool, not a replacement for extension staff.",
    "IPM ordering and severity-based chemical advice support more targeted pesticide use.",
], size=13)
text(s, Inches(5.1), Inches(1.35), Inches(4.4), Inches(0.4), "Limitations", size=17, bold=True, color=ORANGE)
bullets(s, Inches(5.1), Inches(1.85), Inches(4.4), Inches(5), [
    "PlantVillage images are lab photos on plain backgrounds; accuracy on real field photos will be lower until field data is collected.",
    "Only 14 crops / 38 classes; major Indian crops like rice, cotton and chilli need their own image data.",
    "Weather rules are simplified and not yet calibrated with local outbreak records.",
    "Demo field data is synthetic; IoT devices are simulated.",
    "Needs internet for live weather and map tiles.",
], size=13)

# ------------------------------------------------------------------ 18. Future work & conclusion
n += 1
s = content_slide("Future Work & Conclusion", "Discussion", n)
text(s, Inches(0.5), Inches(1.35), Inches(4.4), Inches(0.4), "Future work", size=17, bold=True, color=GREEN)
bullets(s, Inches(0.5), Inches(1.85), Inches(4.4), Inches(5), [
    "Collect field images for rice, cotton, chilli, pulses; fine-tune the full network.",
    "Offline Android app with on-device TensorFlow Lite model.",
    "SMS / WhatsApp / IVR voice alerts for farmers without smartphones.",
    "Calibrate risk models with state pest-surveillance data; add soil-health cards.",
    "Real LoRa/GSM smart traps with automatic insect counting.",
    "Satellite (NDVI) crop-stress layers on the hotspot map.",
], size=13)
rect(s, Inches(5.1), Inches(1.35), Inches(4.4), Inches(5.4), LGREEN)
text(s, Inches(5.3), Inches(1.5), Inches(4.0), Inches(0.4), "Conclusion", size=17, bold=True, color=GREEN)
text(s, Inches(5.3), Inches(2.0), Inches(4.0), Inches(4.6), [
    "KrishiRakshak delivers a complete crop-health system that runs in Chrome: image-based diagnosis, weather-based risk "
    "forecasting, pest-trap and sensor monitoring, hotspot mapping, expert validation, multilingual IPM advice, follow-up and "
    "continual learning, with dashboards for officials.",
    "",
    "It supports the expected outcomes of the problem statement: earlier detection, reduced crop loss, more targeted pesticide use, "
    "faster extension response, wider surveillance coverage and better planning of preventive interventions."], size=13)

# ------------------------------------------------------------------ 19. References / thank you
n += 1
s = content_slide("References", "Discussion", n)
bullets(s, Inches(0.5), Inches(1.4), Inches(9), Inches(4.2), [
    "Hughes, D. P. & Salathé, M. (2015). An open access repository of images on plant health (PlantVillage). arXiv:1511.08060.",
    "Mohanty, S. P., Hughes, D. P. & Salathé, M. (2016). Using deep learning for image-based plant disease detection. Frontiers in Plant Science, 7:1419.",
    "Sandler, M. et al. (2018). MobileNetV2: Inverted residuals and linear bottlenecks. CVPR.",
    "Ester, M. et al. (1996). A density-based algorithm for discovering clusters (DBSCAN). KDD.",
    "Hutton criteria for potato late blight risk — AHDB / UK Met Office (2017).",
    "Mills, W. D. (1944). Efficient use of sulfur dusts and sprays during rain to control apple scab. Cornell Ext. Bull. 630.",
    "Open-Meteo weather forecast API — open-meteo.com.",
    "FAO (2021). Scientific review of the impact of climate change on plant pests.",
], size=12, gap=4)
text(s, Inches(0.5), Inches(5.9), Inches(9), Inches(0.8), "Thank you", size=34, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

out = os.path.join(HERE, "KrishiRakshak_Presentation.pptx")
prs.save(out)
print("saved", out, len(prs.slides), "slides")
